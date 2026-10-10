# -*- coding: utf-8 -*-
"""An address printed beside its own district and state does not repeat them.

The CSP Certificate prints the address and then "DIST- <district>,
<state>" after it. People write their address ending with the district,
state and pincode, so the line read:

    Samvapur belghat gorakhpur up 273404, DIST- Gorakhpur, Uttar Pradesh

Two things decide whether this is safe, and both are tested here.

**Only the tail is trimmed.** A district at the start of an address is
part of the place name — "Patna City Main Road" in Patna district has to
keep its first word — and position is the only way to tell the two apart.

**Only where the template prints them separately.** The ID Card and the
BCP print the address alone, so for them the tail is the only place the
district and pincode appear at all. Trimming there would destroy data.
"""

from datetime import date

import pytest

from app.models.csp_record import CSPRecord
from app.models.enums import DocumentType
from app.services import template_service
from app.services.csp_mapping import address_without_district_state as trim

# Every address on the real records the day this was written, with what it
# should become. These are the shapes people actually type.
REAL_ADDRESSES = [
    ("Samvapur belghat gorakhpur up 273404", "Gorakhpur", "Uttar Pradesh",
     "Samvapur belghat"),
    ("Navipur gajner road Jaipur Akbarpur Kanpur dehat pin code-209311",
     "Kanpur Dehat", "Uttar Pradesh", "Navipur gajner road Jaipur Akbarpur"),
    ("Ninhaura district kanpur dehat 209115", "Kanpur Dehat", "Uttar Pradesh",
     "Ninhaura"),
    ("chotesagaon, bekarsehar, Saharsa, bihar", "Saharsa", "Bihar",
     "chotesagaon, bekarsehar"),
    ("Narayanpur Kothi Fazilnagar Kushinagar 274401", "Kushinagar",
     "Uttar Pradesh", "Narayanpur Kothi Fazilnagar"),
    ("Ward no 7 Subhash Nagar Rasulabad kanpur dehat 209306", "Kanpur Dehat",
     "Uttar Pradesh", "Ward no 7 Subhash Nagar Rasulabad"),
    # Nothing to trim: left exactly as written.
    ("village , rural, urban", "Sitamarhi", "Bihar", "village , rural, urban"),
]


@pytest.mark.parametrize("address,district,state,expected", REAL_ADDRESSES)
def test_real_addresses_trim_as_expected(address, district, state, expected):
    assert trim(address, district=district, state=state) == expected


# ---- what must NOT be trimmed -------------------------------------------

def test_a_district_at_the_start_is_part_of_the_place_name():
    """The case that makes trimming anywhere but the tail unsafe."""
    assert trim("Patna City Main Road", district="Patna", state="Bihar") == (
        "Patna City Main Road")


def test_a_district_in_the_middle_is_kept():
    assert trim("Near Gorakhpur Station, Mohalla Chowk",
                district="Gorakhpur", state="Uttar Pradesh") == (
        "Near Gorakhpur Station, Mohalla Chowk")


def test_an_address_is_never_emptied():
    """An address that is only its district and pincode is left alone.

    Printing nothing is worse than printing something repeated.
    """
    assert trim("Gorakhpur 273404", district="Gorakhpur",
                state="Uttar Pradesh").strip() != ""
    assert trim("Bihar", district="Saharsa", state="Bihar").strip() != ""


def test_a_number_that_is_not_a_pincode_survives():
    assert "209" in trim("House 209, Mohalla Chowk", district="Saharsa",
                         state="Bihar")


def test_a_blank_address_stays_blank():
    assert trim("", district="Saharsa", state="Bihar") == ""
    assert trim(None, district="Saharsa", state="Bihar") == ""


def test_no_district_or_state_means_only_a_pincode_goes():
    assert trim("Mohalla Chowk 273404") == "Mohalla Chowk"


def test_a_single_word_state_is_not_abbreviated_to_one_letter():
    """'B' for Bihar would match far too much to be safe."""
    assert trim("Chowk B", district="Saharsa", state="Bihar") == "Chowk B"


# ---- which documents it applies to --------------------------------------

def _csp(**overrides):
    values = dict(
        csp_id="ADDR001", csp_code="ADDR001", name="Test Name",
        branch_name="Testbranch", branch_code="BR01",
        address="Samvapur belghat gorakhpur up 273404",
        district="Gorakhpur", state="Uttar Pradesh", mobile="9876543210",
        email="test@example.com", ao="Testao", rbo_name="Testrbo",
        asm_name="Testasm", asm_email="asm@example.com",
        asm_contact="9876500000", certificate_date=date(2026, 3, 4),
    )
    values.update(overrides)
    return CSPRecord(**values)


def test_the_certificate_trims_because_it_prints_them_separately():
    values = {"address": "Samvapur belghat gorakhpur up 273404",
              "district": "Gorakhpur", "state": "Uttar Pradesh"}
    placeholders = {'"Address"': "address", '"District"': "district",
                    '"State"': "state"}
    out = template_service._trim_repeated_address(values, placeholders)
    assert out["address"] == "Samvapur belghat"


def test_a_document_printing_only_the_address_keeps_all_of_it():
    """The ID Card and the BCP. Trimming here would lose the district."""
    values = {"address": "Samvapur belghat gorakhpur up 273404",
              "district": "Gorakhpur", "state": "Uttar Pradesh"}
    placeholders = {'"Address"': "address", '"CSP Name"': "csp_name"}
    out = template_service._trim_repeated_address(values, placeholders)
    assert out["address"] == "Samvapur belghat gorakhpur up 273404"


@pytest.mark.parametrize("doc_type,should_trim", [
    (DocumentType.CSP_CERTIFICATE, True),
    (DocumentType.CSP_ID_CARD, False),
    (DocumentType.BCP, False),
])
def test_the_supplied_templates_behave_as_intended(doc_type, should_trim,
                                                   tmp_path):
    """Read back out of the generated file, not asserted on the input."""
    import docx

    from app.services import csp_mapping

    csp = _csp()
    values = csp_mapping.document_values(csp)
    out = str(tmp_path / f"{doc_type.value}.docx")
    template_service.fill_template(doc_type, values, out,
                                   photo_path=template_service.PLACEHOLDER_PHOTO
                                   if doc_type in template_service.PHOTO_MARKERS
                                   else None)

    text = "\n".join("".join(r.text for r in p.runs)
                     for p in template_service._all_paragraphs(
                         docx.Document(out)))

    if should_trim:
        # Title-cased on the way out, which is a separate rule.
        assert "Samvapur Belghat" in text
        assert "273404" not in text, "the pincode was repeated"
        assert "gorakhpur up" not in text.lower(), (
            "the district and state were repeated inside the address")
    else:
        assert "273404" in text, (
            "this document prints only the address, so the pincode is the "
            "only copy there is")


# ---- how values are capitalised -----------------------------------------
#
# Names and places are typed into the sheet however the person felt:
# 'ANIRUDH KUMAR KUSHWAHA', 'jayprakash', 'Mukesh kumar singhi'. On a
# document they should all read the same way. Codes, numbers and email
# addresses are left exactly as entered.

from app.services.csp_mapping import TITLE_CASED_FIELDS, title_case  # noqa: E402


@pytest.mark.parametrize("raw,expected", [
    ("ANIRUDH KUMAR KUSHWAHA", "Anirudh Kumar Kushwaha"),
    ("jayprakash", "Jayprakash"),
    ("Mukesh kumar singhi", "Mukesh Kumar Singhi"),
    ("BELSAND", "Belsand"),
    ("MUZAFFARPUR", "Muzaffarpur"),
    ("chotesagaon, bekarsehar", "Chotesagaon, Bekarsehar"),
    ("Uttar Pradesh", "Uttar Pradesh"),
    ("", ""),
])
def test_values_are_capitalised_word_by_word(raw, expected):
    assert title_case(raw) == expected


@pytest.mark.parametrize("raw", ["Gorakhpur 2", "Khushi Nagar 8", "274401"])
def test_digits_are_left_alone(raw):
    assert title_case(raw) == raw


def test_an_apostrophe_does_not_start_a_new_word():
    """str.title() turns "don't" into "Don'T"; this must not."""
    assert title_case("d'souza colony") == "D'souza Colony"


def test_codes_numbers_and_emails_are_not_capitalised():
    """A CSP code is an identifier and an email address is case-sensitive
    to its owner. Neither is a name."""
    for field in ("csp_code", "branch_code", "mobile", "email", "asm_email",
                  "asm_contact", "photo"):
        assert field not in TITLE_CASED_FIELDS, (
            f"{field} would be rewritten, and it is not a name")


def test_the_printed_values_are_capitalised_but_the_record_is_not():
    """Only the output changes. Nothing is rewritten in the record."""
    from app.services import csp_mapping

    csp = _csp(name="ANIRUDH KUMAR KUSHWAHA", branch_name="BELSAND",
               email="Someone@Example.COM", csp_code="1a851620")
    values = csp_mapping.document_values(csp)

    assert values["csp_name"] == "Anirudh Kumar Kushwaha"
    assert values["branch_name"] == "Belsand"
    # Untouched.
    assert values["email"] == "Someone@Example.COM"
    assert values["csp_code"] == "1a851620"
    # And the record itself still says what was typed.
    assert csp.name == "ANIRUDH KUMAR KUSHWAHA"


def test_the_validator_expects_what_the_filler_produced(tmp_path):
    """The two used to disagree: the certificate's address was trimmed on
    the way in and not on the way out, so validation hunted for text the
    document was never going to contain."""
    from app.services import csp_mapping

    csp = _csp()
    for doc_type in (DocumentType.CSP_CERTIFICATE, DocumentType.CSP_ID_CARD):
        expected = template_service.values_for(
            doc_type, csp_mapping.document_values(csp))
        out = str(tmp_path / f"v_{doc_type.value}.docx")
        template_service.fill_template(
            doc_type, csp_mapping.document_values(csp), out,
            photo_path=template_service.PLACEHOLDER_PHOTO
            if doc_type in template_service.PHOTO_MARKERS else None)

        import docx
        text = "\n".join("".join(r.text for r in p.runs)
                         for p in template_service._all_paragraphs(
                             docx.Document(out)))
        for field in set(
                template_service.placeholders_for(doc_type).values()):
            value = str(expected.get(field, "")).strip()
            if value:
                assert value in text, (
                    f"{doc_type.value}: validation expects {value!r} but the "
                    f"document does not contain it")
