# -*- coding: utf-8 -*-
"""How the CSP is addressed, and whether the result looks like one document.

Two things a reader notices immediately and neither of which any other test
covered: a certificate addressing a woman as "Shri", and a name printed in a
different size or typeface from the sentence around it.

The second was real. The CSP Certificate's "CSP Name" placeholder carries no
size at all, so the name printed at the 11pt document default inside a 22pt
sentence; and in the Training Certificate the name's runs carry no typeface,
so they fell back to the theme font (Aptos) beside Calibri prose.
"""

import os
import tempfile

import pytest
from docx import Document

from app.models.enums import DocumentType
from app.services import template_service as t

pytestmark = pytest.mark.asyncio(loop_scope="session")

PHOTO = os.path.join("app", "assets", "placeholder_photo.png")

VALUES = {
    "csp_name": "Abhishek Anand", "csp_code": "1A850466",
    "branch_name": "Dhabauli", "branch_code": "8567",
    "address": "chotesagaon, Saharsa, Bihar", "mobile": "9955776279",
    "email": "a@example.com", "district": "Saharsa", "state": "Bihar",
    "ao": "Bhagalpur", "rbo_name": "Bhowesh Bhawan",
    "asm_name": "Munna Pathak", "asm_email": "munna.pathak@eko.co.in",
    "asm_contact": "9730953305", "certificate_date": "02.10.2026",
    "certificate_month_year": "October 2026",
}


def _build(doc_type, **overrides):
    values = {**VALUES, **overrides}
    out = os.path.join(tempfile.mkdtemp(), f"{doc_type.value}.docx")
    t.fill_template(
        doc_type, values, out,
        photo_path=PHOTO if doc_type in t.PHOTO_MARKERS else None)
    return out


def _text(path):
    doc = Document(path)
    return "\n".join("".join(r.text for r in p.runs)
                     for p in t._all_paragraphs(doc))


# ---- the salutation ------------------------------------------------------

async def test_a_man_is_addressed_as_mr():
    assert "Mr. Abhishek Anand" in _text(
        _build(DocumentType.TRAINING_CERTIFICATE, gender="Male"))


async def test_a_woman_is_not_addressed_as_shri():
    text = _text(_build(DocumentType.TRAINING_CERTIFICATE, gender="Female"))
    assert "Mrs./Miss" in text
    assert "Shri" not in text


async def test_the_alternatives_are_all_removed():
    """Leaving "Shri/Ms/Mrs." on a certificate reads as an oversight."""
    for gender in ("Male", "Female"):
        text = _text(_build(DocumentType.TRAINING_CERTIFICATE, gender=gender))
        assert t.SALUTATION_TOKEN not in text


async def test_both_occurrences_are_replaced():
    """The Training Certificate says it twice."""
    text = _text(_build(DocumentType.TRAINING_CERTIFICATE, gender="Male"))
    assert text.count("Mr.") >= 2


async def test_an_unknown_gender_keeps_the_general_form():
    """Guessing would address someone wrongly; the template's own text is safe."""
    for gender in ("", "   ", "Other", "Prefer not to say"):
        text = _text(_build(DocumentType.TRAINING_CERTIFICATE, gender=gender))
        assert t.SALUTATION_TOKEN in text


@pytest.mark.parametrize("given,expected", [
    ("Male", "Mr."), ("male", "Mr."), ("  MALE  ", "Mr."),
    ("Female", "Mrs./Miss"), ("FEMALE", "Mrs./Miss"),
    ("", ""), ("Other", ""), (None, ""),
])
async def test_gender_is_read_the_way_the_sheet_and_form_write_it(given, expected):
    assert t.salutation_for(given) == expected


# ---- one document, one typeface -----------------------------------------

@pytest.mark.parametrize("doc_type", list(DocumentType))
async def test_no_paragraph_mixes_sizes_or_typefaces(doc_type):
    """A filled value must look like the sentence it sits in."""
    path = _build(doc_type, gender="Female")
    doc = Document(path)

    for index, para in enumerate(t._all_paragraphs(doc)):
        runs = [r for r in para.runs if r.text.strip()]
        if len(runs) < 2:
            continue
        sizes = {r.font.size.pt if r.font.size else None for r in runs}
        fonts = {r.font.name for r in runs}
        text = "".join(r.text for r in runs)[:60]
        assert len(sizes) == 1, (
            f"{doc_type.value} paragraph {index} mixes sizes {sizes}: {text!r}")
        assert len(fonts) == 1, (
            f"{doc_type.value} paragraph {index} mixes typefaces {fonts}: {text!r}")


async def test_the_csp_name_is_not_left_at_the_document_default():
    """It was printing at 11pt inside a 22pt sentence."""
    doc = Document(_build(DocumentType.CSP_CERTIFICATE, gender="Male"))
    for para in t._all_paragraphs(doc):
        runs = [r for r in para.runs if r.text.strip()]
        if not any("Abhishek Anand" in r.text for r in runs):
            continue
        sizes = {r.font.size.pt for r in runs if r.font.size}
        assert sizes and None not in {r.font.size for r in runs}, (
            "the name has no explicit size and falls back to 11pt")


# ---- the photo marker ----------------------------------------------------

@pytest.mark.parametrize("doc_type", list(t.PHOTO_MARKERS))
async def test_the_registered_photo_marker_is_in_the_template(doc_type):
    """Renaming it in Word must fail loudly, not print the marker on a card."""
    assert t.has_photo_marker(t.template_path(doc_type), doc_type), (
        f"{doc_type.value}: PHOTO_MARKERS says "
        f"{t.PHOTO_MARKERS[doc_type]!r}, which is not in the template")


@pytest.mark.parametrize("doc_type", list(t.PHOTO_MARKERS))
async def test_the_marker_is_gone_from_the_generated_document(doc_type):
    assert not t.has_photo_marker(_build(doc_type, gender="Male"), doc_type)
