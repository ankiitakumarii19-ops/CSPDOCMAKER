"""Taking a CSP's photo from the CSP sheet's Photo column (tracker No. 99).

The sheet held the photos, but the app never read them: the Sheets API
returns a Photo cell as empty, so the ID Card failed with "no usable photo".
The photos are now read from the Excel export of the sheet, by the row they
sit in. Pinned here:

  * each picture goes to the CSP in its own row, byte for byte;
  * a Photo cell with two pictures is not used (whose is it?);
  * a picture outside the Photo column, or in a row with no code, is not used;
  * getting a photo tries the sheet before the CRM;
  * a photo path whose file is gone is missing data, not a photo;
  * a generation failure tells the RM why.
"""

import io
import os

import pytest
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from PIL import Image

from app.services import csp_photo_service, sheet_photo_reader

HEADERS = ["CSP Code", "CSP Name", "Photo"]
PHOTO_CELL = "C"


def _jpeg(colour) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (35, 45), colour).save(buf, "JPEG")
    return buf.getvalue()


def _add_picture(ws, cell, data):
    ws.add_image(XlImage(io.BytesIO(data)), cell)


def _workbook(tabs) -> bytes:
    """tabs: {name: (rows, pictures)}; rows are [code, name]; pictures are
    (cell, bytes)."""
    wb = Workbook()
    wb.remove(wb.active)
    for name, (rows, pictures) in tabs.items():
        ws = wb.create_sheet(name)
        ws.append(HEADERS)
        for row in rows:
            ws.append(list(row))
        for cell, data in pictures:
            _add_picture(ws, cell, data)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


RED, BLUE, GREEN = _jpeg("red"), _jpeg("blue"), _jpeg("green")


def test_each_picture_goes_to_the_csp_in_its_row():
    xlsx = _workbook({"KOData": (
        [("1A000001", "Asha"), ("1A000002", "Babu"), ("1A000003", "Chandu")],
        [(f"{PHOTO_CELL}2", RED), (f"{PHOTO_CELL}4", BLUE)])})
    found = sheet_photo_reader.photos_in_workbook(xlsx, ["KOData"])
    assert set(found.pictures) == {"1A000001", "1A000003"}
    assert found.pictures["1A000001"].data == RED
    assert found.pictures["1A000003"].data == BLUE
    assert (found.pictures["1A000003"].tab, found.pictures["1A000003"].row) == ("KOData", 4)
    # Row 3 is in the sheet, with no picture.
    assert "1A000002" in found.codes


def test_two_pictures_in_one_photo_cell_are_not_used():
    xlsx = _workbook({"KOData": (
        [("1A000001", "Asha"), ("1A000002", "Babu")],
        [(f"{PHOTO_CELL}2", RED), (f"{PHOTO_CELL}2", BLUE),
         (f"{PHOTO_CELL}3", GREEN)])})
    found = sheet_photo_reader.photos_in_workbook(xlsx, ["KOData"])
    assert "1A000001" not in found.pictures
    assert found.ambiguous["1A000001"] == ("KOData", 2)
    assert found.pictures["1A000002"].data == GREEN


def test_a_picture_outside_the_photo_column_or_without_a_code_is_not_used():
    xlsx = _workbook({"KOData": (
        [("1A000001", "Asha"), ("", "No code")],
        [("B2", RED), (f"{PHOTO_CELL}3", BLUE), (f"{PHOTO_CELL}1", GREEN)])})
    found = sheet_photo_reader.photos_in_workbook(xlsx, ["KOData"])
    assert found.pictures == {}


def test_the_sub_ko_tab_is_read_too_and_other_tabs_are_not():
    xlsx = _workbook({
        "KOData": ([("1A000001", "Asha")], [(f"{PHOTO_CELL}2", RED)]),
        "SUBKOData": ([("1A00000101", "Sub")], [(f"{PHOTO_CELL}2", BLUE)]),
        "INFO": ([("1A000009", "Other")], [(f"{PHOTO_CELL}2", GREEN)]),
    })
    found = sheet_photo_reader.photos_in_workbook(xlsx, ["KOData", "SUBKOData"])
    assert set(found.pictures) == {"1A000001", "1A00000101"}


def test_a_tab_without_a_photo_column_gives_no_photos():
    wb = Workbook()
    ws = wb.active
    ws.title = "KOData"
    ws.append(["CSP Code", "CSP Name"])
    ws.append(["1A000001", "Asha"])
    _add_picture(ws, "C2", RED)
    out = io.BytesIO()
    wb.save(out)
    found = sheet_photo_reader.photos_in_workbook(out.getvalue(), ["KOData"])
    assert found.pictures == {}


# ---- fetch_photo: saving it, and saying why when there is none -----------

class _Csp:
    def __init__(self, csp_id, photo_url=None):
        self.csp_id = self.csp_code = csp_id
        self.photo_url = photo_url


@pytest.fixture
def photo_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(csp_photo_service, "photo_directory", lambda: str(tmp_path))
    return tmp_path


def _sheet_with(monkeypatch, tabs):
    found = sheet_photo_reader.photos_in_workbook(_workbook(tabs), list(tabs))
    monkeypatch.setattr(sheet_photo_reader, "sheet_photos",
                        lambda refresh=False: found)


def test_the_sheet_photo_is_saved_unchanged(monkeypatch, photo_dir):
    _sheet_with(monkeypatch, {"KOData": ([("1A000001", "Asha")],
                                         [(f"{PHOTO_CELL}2", RED)])})
    result = sheet_photo_reader.fetch_photo(_Csp("1A000001"))
    assert result.ok, result.message
    assert os.path.dirname(result.path) == str(photo_dir)
    assert open(result.path, "rb").read() == RED
    assert "KOData row 2" in result.message


@pytest.mark.parametrize("code, reason", [
    ("1A000002", sheet_photo_reader.NO_PICTURE),
    ("1A999999", sheet_photo_reader.NOT_IN_SHEET),
])
def test_no_photo_says_why(monkeypatch, photo_dir, code, reason):
    _sheet_with(monkeypatch, {"KOData": (
        [("1A000001", "Asha"), ("1A000002", "Babu")],
        [(f"{PHOTO_CELL}2", RED)])})
    result = sheet_photo_reader.fetch_photo(_Csp(code))
    assert not result.ok
    assert result.reason == reason
    assert result.message == sheet_photo_reader.REASON_TEXT[reason]
    assert os.listdir(photo_dir) == []


def test_a_sheet_that_cannot_be_read_is_reported_not_raised(monkeypatch, photo_dir):
    def broken(refresh=False):
        raise sheet_photo_reader.SheetPhotoError("The CSP sheet could not be read. HttpError: 403")
    monkeypatch.setattr(sheet_photo_reader, "sheet_photos", broken)
    result = sheet_photo_reader.fetch_photo(_Csp("1A000001"))
    assert not result.ok
    assert result.reason == sheet_photo_reader.SHEET_ERROR
    assert "403" in result.message


def test_a_broken_picture_is_not_saved(monkeypatch, photo_dir):
    found = sheet_photo_reader.SheetPhotos(
        pictures={"1A000001": sheet_photo_reader.SheetPicture(
            data=b"not an image", extension=".jpg", tab="KOData", row=2)},
        ambiguous={}, codes={"1A000001"})
    monkeypatch.setattr(sheet_photo_reader, "sheet_photos", lambda refresh=False: found)
    result = sheet_photo_reader.fetch_photo(_Csp("1A000001"))
    assert not result.ok and "not a readable image" in result.message
    assert os.listdir(photo_dir) == []


def test_the_export_is_reused_for_a_minute(monkeypatch):
    """One download serves a whole run of CSPs."""
    monkeypatch.undo()          # the real sheet_photos, not conftest's stub
    reader = sheet_photo_reader
    try:
        calls = []
        xlsx = _workbook({"KOData": ([("1A000001", "Asha")],
                                     [(f"{PHOTO_CELL}2", RED)])})
        monkeypatch.setattr(reader, "_export_workbook",
                            lambda sheet_id: calls.append(sheet_id) or xlsx)
        monkeypatch.setattr(reader.get_settings(), "google_sheets_csp_sheet_id", "SHEET")
        monkeypatch.setattr(reader.get_settings(), "google_sheets_csp_sheet_name", "KOData")
        reader.clear_cache()
        assert "1A000001" in reader.sheet_photos().pictures
        reader.sheet_photos()
        assert calls == ["SHEET"]
        reader.sheet_photos(refresh=True)
        assert len(calls) == 2
    finally:
        reader.clear_cache()


# ---- the order: the sheet, then the CRM ---------------------------------

def test_the_sheet_is_tried_before_the_crm(monkeypatch, photo_dir):
    from app.workers.new_csp_worker import _try_photo
    _sheet_with(monkeypatch, {"KOData": ([("1A000001", "Asha")],
                                         [(f"{PHOTO_CELL}2", RED)])})
    monkeypatch.setattr(csp_photo_service, "fetch_photo",
                        lambda csp: pytest.fail("the CRM must not be asked"))
    csp = _Csp("1A000001")
    assert _try_photo(csp) is True
    assert open(csp.photo_url, "rb").read() == RED


def test_the_crm_is_tried_when_the_sheet_has_no_photo(monkeypatch, photo_dir):
    from app.workers.new_csp_worker import _try_photo
    asked = []

    def crm(csp):
        asked.append(csp.csp_id)
        return csp_photo_service.PhotoResult(False, reason=csp_photo_service.NO_RECORD)
    monkeypatch.setattr(csp_photo_service, "fetch_photo", crm)
    csp = _Csp("1A000001")
    assert _try_photo(csp) is False
    assert asked == ["1A000001"]
    assert _try_photo(csp, from_crm=False) is False
    assert asked == ["1A000001"]


# ---- a photo path whose file is gone is no photo -------------------------

def test_a_photo_path_without_a_file_is_missing_data(monkeypatch, photo_dir):
    from app.config import get_settings
    from app.models import CSPRecord
    from app.models.enums import DocumentType
    from app.services.missing_data_service import check_missing_data

    monkeypatch.setattr(get_settings(), "photo_placeholder_allowed", False)
    csp = CSPRecord(csp_id="1A000001", csp_code="1A000001", name="Asha",
                    photo_url=os.path.join(str(photo_dir), "gone.jpg"))
    labels = {f.field_label for d in check_missing_data(
        csp, [DocumentType.CSP_ID_CARD]).documents for f in d.missing_fields}
    assert "Photo" in labels

    present = photo_dir / "1A000001.jpg"
    present.write_bytes(RED)
    csp.photo_url = str(present)
    labels = {f.field_label for d in check_missing_data(
        csp, [DocumentType.CSP_ID_CARD]).documents for f in d.missing_fields}
    assert "Photo" not in labels


def test_the_csp_is_asked_for_a_photo_whose_file_is_gone(photo_dir):
    from app.services import field_ownership as fo
    csp = _Csp("1A000001", photo_url=os.path.join(str(photo_dir), "gone.jpg"))
    assert "photo_url" in fo.missing(csp)[fo.CSP]
    (photo_dir / "1A000001.jpg").write_bytes(RED)
    assert "photo_url" not in fo.missing(csp)[fo.CSP]
